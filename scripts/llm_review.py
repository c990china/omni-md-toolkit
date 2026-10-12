# -*- coding: utf-8 -*-
"""语义核对作业单：第一段（规则整理）之后，把「需要按原文核对」的点列成清单，
交给**当前模型**拿着原文逐条核对 —— 理解语义、补回真丢的内容、按原表修表格。

分工（与全包口径一致）：
  - **第一段**：`md_reorganizer.py` 纯规则重排版式（不改内容）。
  - **第二段（本脚本 + 当前模型）**：脚本只做**比对与列清单**，不做任何语义判断、
    不调用任何模型、不联网；由**正在使用本技能的模型**对照原文逐条处理。

★ 本段之后**不要再跑 `md_reorganizer`**：`reorganize()` **不幂等** —— 同一文件跑 1 次得 `1.第一项`、
  跑 2 次变 `1. 第一项`、跑 3 次又变回来（周期 2 振荡），多跑一次会让结果漂移。
  核对完直接用 `md_lint.py --source "<源文件>"` 复检即可。

本脚本**不落盘**：作业单默认只打印到 stderr（避免在交付目录多出文件）；
需要留档时用 `--out-list <文件>` 显式指定。

返回码：0 = 未发现待核对项；2 = 有待核对项（待当前模型接手，**不是失败**）；1 = 出错。

用法:
  python scripts/llm_review.py "<源文件>" [--md "<已整理.md>"]
  python scripts/llm_review.py "<源文件>" --out-list "<清单.txt>"
"""
import os
import re
import sys
import argparse
import importlib.util

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
try:
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
MIN_BLOCK = 12          # 少于这么多字的原文块不参与「存在性」比对（噪声太大）


def _find_script(name):
    for f in os.listdir(HERE):
        if f.endswith(name + ".py") or f.endswith("_" + name + ".py"):
            return os.path.join(HERE, f)
    return os.path.join(HERE, name + ".py")


def _load(name):
    target = _find_script(name)
    spec = importlib.util.spec_from_file_location(name, target)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ================= 读原文（**按块**，保留段落边界）=================

def read_blocks(path):
    """返回 (blocks, err)。blocks 为「原文块」列表（尽量一段一块）。

    ★ 不复用 md_lint.source_text()：它对 docx 是**不分段**拼接的，
      无法做段落级核对，故这里按格式各自取块。
    """
    ext = os.path.splitext(path)[1].lower()

    if ext in (".txt", ".md", ".text"):
        with open(path, encoding="utf-8", errors="ignore") as f:
            raw = f.read()
        return [b.strip() for b in re.split(r"\n\s*\n", raw) if b.strip()], None

    if ext == ".docx":
        try:
            import docx
            from docx.oxml.ns import qn
        except Exception:
            return [], "缺 python-docx，无法按段读取 docx"
        d = docx.Document(path)
        blocks = []
        for p in d.element.body.iter(qn("w:p")):
            seg = []
            for node in p.iter():
                if node.tag == qn("w:t"):
                    seg.append(node.text or "")
                elif node.tag in (qn("w:tab"), qn("w:br")):
                    seg.append(" ")
            t = "".join(seg).strip()
            if t:
                blocks.append(t)
        return blocks, None

    if ext == ".pdf":
        try:
            import fitz
        except Exception:
            return [], "缺 PyMuPDF(fitz)，无法按段读取 pdf"
        doc = fitz.open(path)
        try:
            blocks = []
            for i in range(doc.page_count):
                txt = doc.load_page(i).get_text() or ""
                for para in re.split(r"\n\s*\n", txt):
                    for line in para.splitlines():
                        line = line.strip()
                        if line:
                            blocks.append(line)
            return blocks, None
        finally:
            doc.close()

    if ext == ".pptx":
        try:
            from pptx import Presentation
        except Exception:
            return [], "缺 python-pptx，无法按段读取 pptx"
        blocks = []
        for slide in Presentation(path).slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    for para in shape.text_frame.paragraphs:
                        t = "".join(r.text for r in para.runs).strip()
                        if t:
                            blocks.append(t)
        return blocks, None

    # 其余格式（xlsx/csv/html/msg/epub/rtf/odt/doc）：交给 md_lint 的通用读取，
    # 但它是**不分段**的，故按行切块（精度较低，作业单里会注明）。
    try:
        ml = _load("md_lint")
        txt, err = ml.source_text(path)
    except Exception as e:
        return [], "读取源文件失败：%s" % e
    if err:
        return [], err
    if txt is None:
        return [], "不支持的源文件类型：%s" % ext
    return [l.strip() for l in txt.splitlines() if l.strip()], "通用读取（未按段落切分，精度较低）"


# ================= 各项核对 =================

def norm(s, ml):
    return ml.norm_for_loss(s)


def check_missing(blocks, md_norm, ml, limit):
    """原文里有、md 里找不到的块（疑似真丢内容）。"""
    out = []
    for b in blocks:
        nb = norm(b, ml)
        if len(nb) < MIN_BLOCK:
            continue
        if nb not in md_norm:
            out.append(b)
    return out[:limit], len(out)


def check_tables(md_text):
    """md 里的表格问题：缺分隔行 / 列数不一致。"""
    lines = md_text.split("\n")
    blocks, cur = [], []
    for i, ln in enumerate(lines, 1):
        if ln.lstrip().startswith("|"):
            cur.append((i, ln))
        else:
            if cur:
                blocks.append(cur)
                cur = []
    if cur:
        blocks.append(cur)

    problems = []
    for blk in blocks:
        rows = [(n, [c.strip() for c in ln.strip().strip("|").split("|")]) for n, ln in blk]
        if len(rows) < 2:
            problems.append(("表格块只有 %d 行，疑似被压成一行" % len(rows), blk[0][0]))
            continue
        sep = [n for n, cells in rows if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != "")]
        if not sep:
            problems.append(("缺少 `| --- |` 分隔行", rows[0][0]))
        ncols = len(rows[0][1])
        bad = [n for n, cells in rows[1:] if len(cells) != ncols]
        if bad:
            problems.append(("列数不一致（表头 %d 列，第 %s 行列数不同）" % (ncols, bad[:5]), rows[0][0]))
    return problems


def check_headings(md_text):
    """标题骨架问题。"""
    heads = []
    for i, ln in enumerate(md_text.split("\n"), 1):
        m = re.match(r"^(#{1,6})\s+(.*)$", ln.strip())
        if m:
            heads.append((i, len(m.group(1)), m.group(2).strip()))
    problems = []
    h1 = [h for h in heads if h[1] == 1]
    if not h1:
        problems.append(("缺少一级标题（H1）", 0))
    elif len(h1) > 1:
        problems.append(("H1 不唯一，共 %d 个（第 %s 行）" % (len(h1), [h[0] for h in h1][:5]), h1[1][0]))
    prev = None
    for n, lv, txt in heads:
        if prev is not None and lv > prev + 1:
            problems.append(("标题层级跳级：H%d 直接跳到 H%d" % (prev, lv), n))
        prev = lv
    for n, lv, txt in heads:
        if txt and txt[-1] in "。！？；，、":
            problems.append(("标题以标点结尾（疑似假标题）：「%s」" % txt[:30], n))
    return problems


def check_images(md_text, md_path):
    """图片引用断链 / alt 缺失。"""
    base = os.path.dirname(os.path.abspath(md_path))
    broken, noalt = [], []
    for i, ln in enumerate(md_text.split("\n"), 1):
        for m in re.finditer(r"!\[([^\]]*)\]\(([^)\s]+)\)", ln):
            alt, src = m.group(1).strip(), m.group(2).strip()
            if src.startswith(("http://", "https://")):
                continue
            p = os.path.normpath(os.path.join(base, src.replace("/", os.sep)))
            if not os.path.exists(p):
                broken.append((i, src))
            if (not alt) or alt == os.path.splitext(os.path.basename(src))[0]:
                noalt.append((i, src))
    return broken, noalt


# ================= 作业单 =================

def build_handoff(src, md_path, missing, n_missing, tables, heads, broken, noalt,
                  block_note, max_each):
    L = []
    L.append("[语义核对] 第一段（规则整理）已完成，下面是**需要按原文核对**的点")
    L.append("  原文: %s" % os.path.abspath(src))
    L.append("  已整理 MD: %s" % os.path.abspath(md_path))
    L.append("")
    L.append("  ▶ 请由**你（当前模型）**打开原文（必要时把 PDF/图片页面渲染出来看图），逐条核对后")
    L.append("     用你的理解修正上面这个 md；本脚本**不做语义判断、不调用任何模型**。")
    if block_note:
        L.append("  （注：%s）" % block_note)
    L.append("")

    L.append("  ① 疑似缺失的原文内容：%d 处" % n_missing)
    if missing:
        for b in missing:
            L.append("     - %s" % b[:100].replace("\n", " "))
        if n_missing > len(missing):
            L.append("     ...（共 %d 处，此处仅列前 %d）" % (n_missing, len(missing)))
        L.append("     → 请对照原文确认：是**真丢了**（补回），还是仅被换行/改写（无需处理）。")
    else:
        L.append("     （未发现整段缺失）")
    L.append("")

    L.append("  ② 表格问题：%d 处" % len(tables))
    for msg, ln in tables[:max_each]:
        L.append("     - 第 %s 行：%s" % (ln, msg))
    if not tables:
        L.append("     （表格结构正常）")
    elif tables:
        L.append("     → 请按**原文表格**重建；也可先用 fix_tables.py 做规则级修复。")
    L.append("")

    L.append("  ③ 标题骨架问题：%d 处" % len(heads))
    for msg, ln in heads[:max_each]:
        L.append("     - %s：%s" % (("第 %s 行" % ln) if ln else "全文", msg))
    if not heads:
        L.append("     （标题骨架正常）")
    L.append("")

    L.append("  ④ 图片：断链 %d 处；alt 缺失/占位 %d 处" % (len(broken), len(noalt)))
    for ln, s in broken[:max_each]:
        L.append("     - 断链 第 %s 行：%s" % (ln, s))
    for ln, s in noalt[:max_each]:
        L.append("     - alt 占位 第 %s 行：%s" % (ln, s))
    if noalt:
        L.append("     → alt 可用 media_extract --recognize 出作业单后由你写回。")
    if not broken and not noalt:
        L.append("     （图片引用正常）")
    L.append("")
    L.append("  ★ 修正原则：**以原文为准**。原文没有的，不要凭印象添加；原文有的，补回原样。")
    L.append("     改完可用 md_lint.py --source \"<源文件>\" 复检。")
    return "\n".join(L)


def main():
    ap = argparse.ArgumentParser(
        description="语义核对作业单：列出「需按原文核对」的点，交给当前模型处理（不调模型、不联网）")
    ap.add_argument("src", help="源文件（docx/pdf/pptx/txt/md/xlsx/csv/html…）")
    ap.add_argument("--md", default=None, help="已整理的 MD（默认取与源文件同名的 <基名>.md）")
    ap.add_argument("--out-list", default=None, help="把作业单另存为该文件（默认只打印，不落盘）")
    ap.add_argument("--max", type=int, default=30, help="每类最多列几条（默认 30）")
    a = ap.parse_args()

    src = os.path.abspath(a.src)
    if not os.path.exists(src):
        sys.stderr.write("[错误] 源文件不存在：%s\n" % src)
        return 1
    md_path = a.md or (os.path.splitext(src)[0] + ".md")
    if not os.path.exists(md_path):
        sys.stderr.write("[错误] 找不到已整理 MD：%s\n" % md_path)
        sys.stderr.write("[提示] 请先跑 convert.py / pipeline.py 完成**第一段（规则整理）**。\n")
        return 1

    try:
        ml = _load("md_lint")
    except Exception as e:
        sys.stderr.write("[错误] 无法加载 md_lint.py（用于归一化与比对）：%s\n" % e)
        return 1

    with open(md_path, encoding="utf-8", errors="ignore") as f:
        md_text = f.read()
    md_norm = norm(md_text, ml)

    blocks, block_note = read_blocks(src)
    if not blocks:
        sys.stderr.write("[错误] 读不到原文内容：%s\n" % (block_note or "未知原因"))
        return 1

    missing, n_missing = check_missing(blocks, md_norm, ml, a.max)
    tables = check_tables(md_text)
    heads = check_headings(md_text)
    broken, noalt = check_images(md_text, md_path)

    text = build_handoff(src, md_path, missing, n_missing, tables, heads, broken, noalt,
                         block_note, a.max)

    if a.out_list:
        with open(a.out_list, "w", encoding="utf-8", newline="") as f:
            f.write(text + "\n")
        sys.stderr.write("[核对] 作业单已写入：%s\n" % os.path.abspath(a.out_list))
    sys.stderr.write(text + "\n")

    total = n_missing + len(tables) + len(heads) + len(broken) + len(noalt)
    if total == 0:
        sys.stderr.write("[核对] 未发现待核对项。\n")
        return 0
    sys.stderr.write("[核对] 共 %d 处待核对 —— 请由当前模型按原文处理。\n" % total)
    return 2          # 2 = 待当前模型接手（不是失败）


if __name__ == "__main__":
    sys.exit(main())
